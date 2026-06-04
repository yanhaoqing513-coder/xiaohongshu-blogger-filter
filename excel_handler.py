"""
Excel处理模块 - 读取和写入Excel文件
"""
import os
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils.cell import coordinate_to_tuple
from datetime import datetime


FIXED_COLUMNS = {
    "nickname": 1,
    "link": 2,
    "screenshot": 3,
    "matched": 4,
    "reason": 5,
    "note": 6,
}

FIXED_HEADERS = ["昵称", "主页链接", "主页截图", "是否合适", "理由", "备注"]


def _image_starts_at(image, row: int, column: int) -> bool:
    """判断工作表图片是否锚定在指定单元格。"""
    anchor = getattr(image, "anchor", None)
    if isinstance(anchor, str):
        try:
            anchor_row, anchor_col = coordinate_to_tuple(anchor)
            return anchor_row == row and anchor_col == column
        except ValueError:
            return False

    anchor_from = getattr(anchor, "_from", None)
    if anchor_from is None:
        return False

    return anchor_from.row == row - 1 and anchor_from.col == column - 1


def get_cell_link(cell) -> str:
    """从单元格提取链接（支持文本和超链接）"""
    # 1. 检查超链接
    if cell.hyperlink and cell.hyperlink.target:
        try:
            target = cell.hyperlink.target
            if "xiaohongshu.com" in str(target):
                return str(target).strip()
        except:
            pass
            
    # 2. 检查单元格文本
    if cell.value and "xiaohongshu.com" in str(cell.value):
        return str(cell.value).strip()
    return None


def read_excel_links(file_path: str) -> list[dict]:
    """
    读取Excel文件中的链接列表
    
    Args:
        file_path: Excel文件路径
        
    Returns:
        包含链接信息的字典列表 [{"row": 行号, "link": 链接}, ...]
    """
    workbook = load_workbook(file_path)
    sheet = workbook.active
    
    links = []
    link_col = None
    
    # 查找包含链接的列（查找表头或直接识别URL）
    for col in range(1, sheet.max_column + 1):
        # 1. 检查表头
        header = sheet.cell(row=1, column=col).value
        if header and any(keyword in str(header).lower() for keyword in ["链接", "link", "url", "主页"]):
            link_col = col
            break
            
        # 2. 检查第一行（如果没有表头，第一行可能是数据）
        first_cell = sheet.cell(row=1, column=col)
        if get_cell_link(first_cell):
            link_col = col
            break
            
        # 3. 检查第二行（如果有表头，第二行是数据）
        second_cell = sheet.cell(row=2, column=col)
        if get_cell_link(second_cell):
            link_col = col
            break
    
    if link_col is None:
        # 默认使用第一列
        link_col = 1
    
    # 读取链接数据
    # 如果第一行看起来像链接，从第一行开始；否则从第二行开始
    first_cell = sheet.cell(row=1, column=link_col)
    start_row = 1 if get_cell_link(first_cell) else 2
    
    for row in range(start_row, sheet.max_row + 1):
        cell = sheet.cell(row=row, column=link_col)
        link = get_cell_link(cell)
        
        if link:
            links.append({
                "row": row,
                "link": link
            })
    
    workbook.close()
    return links


def read_fixed_format_links(file_path: str) -> list[dict]:
    """
    读取固定格式Excel：
    昵称、主页链接、主页截图、是否合适、理由、备注
    已经填写“是否合适”的行会跳过，用于断点续跑。
    """
    workbook = load_workbook(file_path)
    sheet = workbook.active
    links = []

    for row in range(1, sheet.max_row + 1):
        link_cell = sheet.cell(row=row, column=FIXED_COLUMNS["link"])
        link = get_cell_link(link_cell)
        if not link:
            continue

        matched_value = sheet.cell(row=row, column=FIXED_COLUMNS["matched"]).value
        if matched_value not in (None, ""):
            continue

        links.append({
            "row": row,
            "nickname": sheet.cell(row=row, column=FIXED_COLUMNS["nickname"]).value or "",
            "link": link
        })

    workbook.close()
    return links


def ensure_fixed_format_headers(file_path: str) -> None:
    """确保固定格式表头存在；如果第一行不是表头，不主动插入新行。"""
    workbook = load_workbook(file_path)
    sheet = workbook.active

    existing = [sheet.cell(row=1, column=col).value for col in range(1, 7)]
    if any(value in existing for value in FIXED_HEADERS):
        for col, header in enumerate(FIXED_HEADERS, 1):
            if not sheet.cell(row=1, column=col).value:
                sheet.cell(row=1, column=col, value=header)

    workbook.save(file_path)
    workbook.close()


def write_fixed_format_result(
    file_path: str,
    row: int,
    screenshot_path: str,
    matched: bool,
    reason: str,
    note: str,
) -> None:
    """将单行筛选结果写回固定格式Excel。"""
    workbook = load_workbook(file_path)
    sheet = workbook.active

    sheet.cell(row=row, column=FIXED_COLUMNS["matched"], value="合适" if matched else "不合适")
    sheet.cell(row=row, column=FIXED_COLUMNS["reason"], value=reason or "")
    sheet.cell(row=row, column=FIXED_COLUMNS["note"], value=note or "")

    if screenshot_path and os.path.exists(screenshot_path):
        screenshot_column = FIXED_COLUMNS["screenshot"]
        sheet._images = [
            image for image in sheet._images
            if not _image_starts_at(image, row, screenshot_column)
        ]

        image = ExcelImage(screenshot_path)
        image.width = 180
        image.height = 120
        cell_ref = sheet.cell(row=row, column=screenshot_column).coordinate
        sheet.add_image(image, cell_ref)
        sheet.row_dimensions[row].height = 95

    sheet.column_dimensions["C"].width = 26
    sheet.column_dimensions["D"].width = 12
    sheet.column_dimensions["E"].width = 24
    sheet.column_dimensions["F"].width = 60

    workbook.save(file_path)
    workbook.close()


def create_result_excel(results: list[dict], output_path: str) -> str:
    """
    创建筛选结果Excel文件
    
    Args:
        results: 筛选结果列表
        output_path: 输出文件路径
        
    Returns:
        输出文件路径
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "筛选结果"
    
    # 设置表头
    headers = ["序号", "主页链接", "是否符合标准", "判断理由", "账号分析", "主页内容总结", "分析方式"]
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = cell.font.copy(bold=True)
    
    # 写入数据
    for idx, result in enumerate(results, 1):
        sheet.cell(row=idx + 1, column=1, value=idx)
        sheet.cell(row=idx + 1, column=2, value=result.get("link", ""))
        sheet.cell(row=idx + 1, column=3, value="符合" if result.get("matched", False) else "不符合")
        sheet.cell(row=idx + 1, column=4, value=result.get("reason", ""))
        sheet.cell(row=idx + 1, column=5, value=result.get("account_analysis", ""))
        # 主页内容总结，限制30字
        summary = result.get("content_summary", "")
        if len(summary) > 30:
            summary = summary[:30]
        sheet.cell(row=idx + 1, column=6, value=summary)
        mode_text = "大模型图片识别" if result.get("analysis_mode") == "vision" else "OCR + 文本识别"
        sheet.cell(row=idx + 1, column=7, value=mode_text)
    
    # 调整列宽
    sheet.column_dimensions['A'].width = 8
    sheet.column_dimensions['B'].width = 50
    sheet.column_dimensions['C'].width = 14
    sheet.column_dimensions['D'].width = 50
    sheet.column_dimensions['E'].width = 60
    sheet.column_dimensions['F'].width = 40
    sheet.column_dimensions['G'].width = 18
    
    workbook.save(output_path)
    workbook.close()
    
    return output_path


def create_collection_excel(results: list[dict], output_path: str) -> str:
    """
    创建数据采集结果Excel文件
    
    Args:
        results: 采集结果列表
        output_path: 输出文件路径
        
    Returns:
        输出文件路径
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "采集结果"
    
    # 设置表头
    headers = ["序号", "主页链接", "粉丝数", "采集状态"]
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = cell.font.copy(bold=True)
    
    # 写入数据
    for idx, result in enumerate(results, 1):
        sheet.cell(row=idx + 1, column=1, value=idx)
        sheet.cell(row=idx + 1, column=2, value=result.get("link", ""))
        sheet.cell(row=idx + 1, column=3, value=result.get("fans", ""))
        sheet.cell(row=idx + 1, column=4, value=result.get("status", ""))
    
    # 调整列宽
    sheet.column_dimensions['A'].width = 8
    sheet.column_dimensions['B'].width = 50
    sheet.column_dimensions['C'].width = 15
    sheet.column_dimensions['D'].width = 15
    
    workbook.save(output_path)
    workbook.close()
    
    return output_path


def generate_output_filename(original_filename: str) -> str:
    """生成输出文件名"""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    name, ext = os.path.splitext(original_filename)
    return f"{name}_筛选结果_{timestamp}.xlsx"
